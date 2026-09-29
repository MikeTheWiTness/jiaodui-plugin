"""jiaodui/docx_format_enhancer.py 单元测试（移植自旧仓 tests/test_docx_format_enhancer.py）。

依赖 python-docx 的用例标 skipif；依赖 pandoc 的用例用 shutil.which 判断 skip。
"""
import importlib.util
import shutil
import subprocess

import pytest

from jiaodui.docx_format_enhancer import (extract_special_formats, inject_format_markers,
                                          strip_format_markers)
from jiaodui.docx_report import find_pandoc

HAS_DOCX = importlib.util.find_spec("docx") is not None
PANDOC = shutil.which("pandoc") or find_pandoc()
REQUIRES_DOCX = pytest.mark.skipif(not HAS_DOCX, reason="python-docx 不可用")
REQUIRES_PANDOC = pytest.mark.skipif(PANDOC is None, reason="pandoc 不可用")


def create_test_docx(output_path):
    """创建一个包含各种格式的测试 Word 文档"""
    from docx import Document
    from docx.enum.text import WD_UNDERLINE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()

    def add_emphasis_mark(run):
        """给 run 加着重号（通过 XML 方式）"""
        rPr = run._element.get_or_add_rPr()
        emph = OxmlElement('w:emph')
        emph.set(qn('w:val'), 'dot')
        rPr.append(emph)

    # 标题
    doc.add_heading('测试特殊格式', level=1)

    # 1. 着重号
    p1 = doc.add_paragraph()
    p1.add_run('这是一段普通文字，')
    run_emph = p1.add_run('这几个字有')
    add_emphasis_mark(run_emph)
    run_emph2 = p1.add_run('着重号')
    add_emphasis_mark(run_emph2)
    p1.add_run('，后面是普通文字。')

    # 2. 下划线
    p2 = doc.add_paragraph()
    p2.add_run('这句话中，')
    run_under = p2.add_run('这里有下划线')
    run_under.font.underline = True
    p2.add_run('，其他没有。')

    # 3. 波浪线
    p3 = doc.add_paragraph()
    p3.add_run('这句话中，')
    run_wavy = p3.add_run('这里有波浪线')
    run_wavy.font.underline = WD_UNDERLINE.WAVY
    p3.add_run('，其他没有。')

    # 4. 删除线
    p4 = doc.add_paragraph()
    p4.add_run('这句话中，')
    run_strike = p4.add_run('这里有删除线')
    run_strike.font.strike = True
    p4.add_run('，其他没有。')

    # 5. 下标和上标
    p5 = doc.add_paragraph()
    p5.add_run('化学公式：H')
    run_sub = p5.add_run('2')
    run_sub.font.subscript = True
    p5.add_run('O，数学：x')
    run_sup = p5.add_run('2')
    run_sup.font.superscript = True
    p5.add_run(' + y')
    run_sup2 = p5.add_run('2')
    run_sup2.font.superscript = True
    p5.add_run(' = z')
    run_sup3 = p5.add_run('2')
    run_sup3.font.superscript = True

    # 6. 粗体和斜体（Pandoc 原生支持）
    p6 = doc.add_paragraph()
    p6.add_run('这句话有')
    run_bold = p6.add_run('粗体')
    run_bold.bold = True
    p6.add_run('和')
    run_italic = p6.add_run('斜体')
    run_italic.italic = True
    p6.add_run('。')

    # 7. 古诗词（语文场景）
    doc.add_heading('古诗词示例', level=2)
    p7 = doc.add_paragraph()
    run_title = p7.add_run('静夜思')
    run_title.bold = True
    p7.add_run('\n')
    p7.add_run('床前明月光，')
    p7.add_run('\n')
    run_emph_poem = p7.add_run('疑是地上霜')
    add_emphasis_mark(run_emph_poem)
    p7.add_run('。')
    p7.add_run('\n')
    p7.add_run('举头望明月，')
    p7.add_run('\n')
    run_emph_poem2 = p7.add_run('低头思故乡')
    run_emph_poem2.font.underline = WD_UNDERLINE.WAVY
    p7.add_run('。')

    doc.save(output_path)


def _docx_to_md(pandoc, docx_path, md_path):
    r = subprocess.run(
        [pandoc, "-f", "docx", "-t", "markdown-smart", "--wrap", "none",
         str(docx_path), "-o", str(md_path)],
        capture_output=True, text=True)
    return r.returncode == 0


@REQUIRES_DOCX
def test_extract_formats(tmp_path):
    """测试提取特殊格式"""
    docx_path = tmp_path / "test.docx"
    create_test_docx(str(docx_path))

    formats = extract_special_formats(str(docx_path))
    assert len(formats) > 0, "应该提取到至少一个格式"
    by_type = {}
    for fmt in formats:
        by_type.setdefault(fmt["type"], []).append(fmt["text"])
    for expected in ("emphasis_dot", "underline", "underline_wavy", "strike"):
        assert expected in by_type, f"缺少格式类型 {expected}: {sorted(by_type)}"
        assert by_type[expected], f"{expected} 无文本"


@REQUIRES_DOCX
@REQUIRES_PANDOC
def test_inject_markers(tmp_path):
    """测试注入格式标记：pandoc 转换后应至少注入一种自定义标记"""
    docx_path = tmp_path / "test.docx"
    create_test_docx(str(docx_path))

    md_path = tmp_path / "test.md"
    assert _docx_to_md(PANDOC, docx_path, md_path), "Pandoc 转换应该成功"
    md_text = md_path.read_text(encoding="utf-8")

    enhanced = inject_format_markers(md_text, str(docx_path))
    markers_found = [m for m in ("<着重>", "<下划线>", "<波浪线>", "<删除线>")
                     if m in enhanced]
    assert markers_found, f"应至少注入一种格式标记；markdown={md_text!r}"


def test_strip_markers():
    """测试移除格式标记"""
    test_text = "这是<着重>测试</着重>文本，有<下划线>下划线</下划线>和<波浪线>波浪线</波浪线>"
    stripped = strip_format_markers(test_text)
    expected = "这是测试文本，有下划线和波浪线"
    assert stripped == expected, f"清理后不匹配: {stripped} != {expected}"
