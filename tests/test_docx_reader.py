"""docx 原生读取与结构拆分的契约测试。

用 python-docx 现场构造固定模板夹具；另用仓库内冻结的 m1 语料验证缺图占位图
过滤。不依赖 Downloads 里的真实材料。
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("PIL")

from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image

from jiaodui.docx_reader import (DocxReadError, UnsupportedTemplate, read_docx)
from jiaodui.errors import ContractError
from jiaodui.split import split_structural

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _build_template(path: Path) -> Path:
    """构造最小固定模板：标题表 + 五链题目表 + 题干内联图 + 解析数据表。"""
    image = path.parent / "tiny.png"
    Image.new("RGB", (40, 40), "white").save(image)

    document = Document()
    document.add_heading("模块一 测试", level=2)

    heading = document.add_table(rows=1, cols=1)
    heading.cell(0, 0).paragraphs[0].style = document.styles["Heading 3"]
    heading.cell(0, 0).paragraphs[0].add_run("必备知识")

    question = document.add_table(rows=5, cols=1)
    title = question.cell(0, 0).paragraphs[0]
    title.add_run("练1").bold = True
    title.add_run("（2026·测试卷）")

    stem = question.cell(1, 0).add_table(rows=1, cols=1)
    stem.cell(0, 0).paragraphs[0].add_run("题干：合力$F$随时间变化如表，求加速度。")
    stem.cell(0, 0).paragraphs[0].add_run().add_picture(str(image), width=Inches(0.5))

    options = question.cell(2, 0).add_table(rows=2, cols=2)
    for row, labels in enumerate((("A", "B"), ("C", "D"))):
        for col, label in enumerate(labels):
            options.cell(row, col).text = f"{label}. 选项{label}"

    question.cell(3, 0).text = ""

    answer = question.cell(4, 0).add_table(rows=3, cols=2)
    answer.cell(0, 0).text = "答案:"
    answer.cell(0, 1).text = "A"
    answer.cell(1, 0).text = "解答："
    answer.cell(1, 1).text = r"由牛顿第二定律得$a=\dfrac{F}{m}$。"
    data = answer.cell(2, 1).add_table(rows=3, cols=3)
    for row, values in enumerate((("次数", "U/V", "I/A"),
                                  ("1", "1", "0.2"),
                                  ("2", "2", "0.4"))):
        for col, value in enumerate(values):
            data.cell(row, col).text = value
    data.style = document.styles["Table Grid"]

    target = path / "fixture.docx"
    document.save(target)
    return target


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return "".join(node.text or "" for node in root.iter(f"{{{W}}}t"))


def _markdown_plain(text: str) -> str:
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\|(?:\s*---\s*\|)+$", "", text, flags=re.MULTILINE)
    text = text.replace("**", "").replace("|", "").replace("<br>", "")
    return re.sub(r"\s+", "", text)


def _normal(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _add_unknown_table(document):
    """两行两列、每行只有一个非空单元格：既不是题目/标题，也不是数据表。"""
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(1, 0).text = "乙"
    return table


def test_reads_template_options_and_data_table(tmp_path):
    source = _build_template(tmp_path)
    media = tmp_path / "media"
    document = read_docx(source, media_dir=media, link_prefix="./fixture_images/media")

    assert document.stats["questions"] == 1
    assert document.stats["headings"] == 2
    assert document.stats["data_tables"] == 1
    assert document.stats["images"] == 1

    text = document.markdown
    assert "## 模块一 测试" in text
    assert "### 必备知识" in text
    assert "**练1**（2026·测试卷）" in text
    # 2×2 选项布局必须展开成四行，不能当数据表
    for line in ("A. 选项A", "B. 选项B", "C. 选项C", "D. 选项D"):
        assert line in text
    assert "答案: A" in text
    assert "| 次数 | U/V | I/A |" in text
    assert "| 1 | 1 | 0.2 |" in text
    assert (media / "image1.png").is_file()


def test_text_is_faithful_to_docx(tmp_path):
    source = _build_template(tmp_path)
    document = read_docx(source, media_dir=tmp_path / "media",
                         link_prefix="./fixture_images/media")
    assert _normal(_docx_text(source)) == _normal(_markdown_plain(document.markdown))


def test_nested_table_inside_single_cell_is_kept(tmp_path):
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "L与d的关系"
    inner.cell(0, 1).text = "图示"
    inner.cell(1, 0).text = "$L<d$"
    inner.cell(1, 1).text = "进"
    inner.style = document.styles["Table Grid"]
    source = tmp_path / "nested.docx"
    document.save(source)

    result = read_docx(source, media_dir=tmp_path / "media")

    assert "| L与d的关系 | 图示 |" in result.markdown
    assert "| $L<d$ | 进 |" in result.markdown


def test_single_row_layout_table_flattens(tmp_path):
    document = Document()
    document.add_heading("模块一 测试", level=2)
    bar = document.add_table(rows=1, cols=3)
    bar.cell(0, 1).text = "模型识别"
    source = tmp_path / "layout.docx"
    document.save(source)

    result = read_docx(source, media_dir=tmp_path / "media")

    assert "模型识别" in result.markdown
    assert result.stats["tables"] == 0


def test_unknown_table_shape_raises_unsupported(tmp_path):
    document = Document()
    _add_unknown_table(document)
    source = tmp_path / "bad.docx"
    document.save(source)

    with pytest.raises(UnsupportedTemplate):
        read_docx(source, media_dir=tmp_path / "media")


def test_broken_image_placeholders_are_skipped(tmp_path):
    from jiaodui.convert import _is_broken_image_placeholder

    fixture = (Path(__file__).resolve().parents[1]
               / "evaluation/m1/source/第 6 讲校对测试.docx")
    if not fixture.is_file():
        pytest.skip("缺少 m1 冻结语料")
    media = tmp_path / "media"

    document = read_docx(fixture, media_dir=media, link_prefix="x")

    assert document.stats["images"] == 5
    assert not any(_is_broken_image_placeholder(path)
                   for path in media.iterdir() if path.is_file())
    assert any("缺图占位图" in warning for warning in document.warnings)
    assert sorted(item["name"] for item in document.images) == [
        "image4.png", "image5.png", "image6.png", "image7.png", "image8.png"]


def test_broken_package_raises_read_error(tmp_path):
    source = tmp_path / "broken.docx"
    source.write_text("not a zip", encoding="utf-8")
    with pytest.raises(DocxReadError):
        read_docx(source, media_dir=tmp_path / "media")


def test_structural_split_keeps_teacher_version_questions(tmp_path, monkeypatch):
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    raw = tmp_path / "材料_raw.md"
    raw.write_text(
        "**练1**\n"
        "题干一\n"
        "\n"
        "**教师版-基础练习**(★★★☆☆)\n"
        "题干二\n"
        "\n"
        "**练2**\n"
        "题干三\n",
        encoding="utf-8")
    blocks = [
        {"kind": "question", "start_line": 1, "end_line": 2, "title": "**练1**"},
        {"kind": "question", "start_line": 4, "end_line": 5,
         "title": "**教师版-基础练习**(★★★☆☆)"},
        {"kind": "question", "start_line": 7, "end_line": 8, "title": "**练2**"},
    ]

    result = split_structural(str(raw), blocks, str(tmp_path / "units"), "材料")

    assert [unit["title"] for unit in result.units] == [
        "**练1**", "**教师版-基础练习**(★★★☆☆)", "**练2**"]
    assert "题干二" in Path(result.units[1]["file"]).read_text(encoding="utf-8")
    assert "题干三" not in Path(result.units[1]["file"]).read_text(encoding="utf-8")


def test_structural_split_splits_knowledge_at_sections(tmp_path, monkeypatch):
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    raw = tmp_path / "材料_raw.md"
    raw.write_text(
        "### 直击课堂\n"
        "![](./images/nav.png)\n"
        "\n"
        "## 模块一\n"
        "### 模型大招\n"
        "模型正文\n"
        "\n"
        "### 必备知识\n"
        "知识正文\n"
        "\n"
        "**练1**\n"
        "题干\n",
        encoding="utf-8")
    blocks = [
        {"kind": "heading", "level": 3, "start_line": 1, "end_line": 1, "title": "直击课堂"},
        {"kind": "paragraph", "level": 0, "start_line": 2, "end_line": 2, "title": ""},
        {"kind": "heading", "level": 2, "start_line": 4, "end_line": 4, "title": "模块一"},
        {"kind": "heading", "level": 3, "start_line": 5, "end_line": 5, "title": "模型大招"},
        {"kind": "paragraph", "level": 0, "start_line": 6, "end_line": 6, "title": ""},
        {"kind": "heading", "level": 3, "start_line": 8, "end_line": 8, "title": "必备知识"},
        {"kind": "paragraph", "level": 0, "start_line": 9, "end_line": 9, "title": ""},
        {"kind": "question", "level": 0, "start_line": 11, "end_line": 12, "title": "**练1**"},
    ]

    result = split_structural(str(raw), blocks, str(tmp_path / "units"), "材料")

    assert [unit["title"] for unit in result.units] == [
        "直击课堂", "模块一", "必备知识", "**练1**"]
    model = Path(result.units[1]["file"]).read_text(encoding="utf-8")
    assert "### 模型大招" in model and "模型正文" in model


def test_structural_split_drops_heading_only_run(tmp_path, monkeypatch):
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    raw = tmp_path / "材料_raw.md"
    raw.write_text("## 模块一\n#### 基础演练\n**练1**\n题干\n", encoding="utf-8")
    blocks = [
        {"kind": "heading", "start_line": 1, "end_line": 1, "title": "模块一"},
        {"kind": "heading", "start_line": 2, "end_line": 2, "title": "基础演练"},
        {"kind": "question", "start_line": 3, "end_line": 4, "title": "**练1**"},
    ]

    result = split_structural(str(raw), blocks, str(tmp_path / "units"), "材料")

    assert len(result.units) == 1
    assert result.units[0]["title"] == "**练1**"


def test_convert_uses_native_engine(tmp_path, monkeypatch):
    from jiaodui.convert import convert_to_raw

    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    source = _build_template(tmp_path)
    result = convert_to_raw(str(source), str(tmp_path / "raw"), "材料", mode="lecture")

    assert result.engine == "docx"
    assert result.blocks
    assert result.copied == 1
    text = Path(result.raw_md).read_text(encoding="utf-8")
    assert "| 次数 | U/V | I/A |" in text


def test_convert_falls_back_to_pandoc(tmp_path, monkeypatch):
    from jiaodui import convert

    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    document = Document()
    _add_unknown_table(document)
    source = tmp_path / "bad.docx"
    document.save(source)

    def fake_pandoc(src, output_md, img_dir, use_mathjax=False):
        Path(output_md).write_text("# pandoc 回退\n", encoding="utf-8")
        return True

    monkeypatch.setattr(convert, "check_pandoc", lambda: True)
    monkeypatch.setattr(convert, "convert_with_pandoc", fake_pandoc)
    result = convert.convert_to_raw(str(source), str(tmp_path / "raw"), "材料",
                                   mode="lecture")

    assert result.engine == "pandoc"
    assert any("回退 pandoc" in warning for warning in result.warnings)
    assert Path(result.raw_md).read_text(encoding="utf-8").startswith("# pandoc")


def test_import_material_returns_compact_summary(tmp_path, monkeypatch):
    from jiaodui.workflow import import_material

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source_dir = tmp_path / "src"
    source_dir.mkdir()
    source = _build_template(source_dir)

    summary = import_material(str(source), subject="高中物理")

    assert summary["ok"] is True
    assert summary["mode"] == "lecture"
    assert summary["mode_origin"] == "structure"
    assert summary["engine"] == "docx"
    assert summary["units"] == 1
    assert summary["question_units"] == 1
    assert summary["knowledge_units"] == 0
    assert summary["copied"] == 1
    assert summary["first_units"][0]["name"] == "单元1"
    # 摘要保持紧凑：不携带完整单元列表与 unit_ids
    assert "unit_ids" not in summary
    assert "units_detail" not in summary


def test_forced_native_engine_rejects_unknown_structure(tmp_path, monkeypatch):
    from jiaodui import convert
    from jiaodui.errors import UnsupportedError

    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    document = Document()
    _add_unknown_table(document)
    source = tmp_path / "bad.docx"
    document.save(source)

    with pytest.raises(UnsupportedError):
        convert.convert_to_raw(str(source), str(tmp_path / "raw"), "材料",
                               mode="lecture", import_engine="docx")


def test_navigation_units_are_marked_skip(tmp_path, monkeypatch):
    from jiaodui.split import mark_navigation_units

    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    root = tmp_path / "units" / "材料"
    for name, first_line in (("单元1", "### 直击课堂"), ("单元2", "### 模型")):
        unit = root / name
        unit.mkdir(parents=True)
        (unit / f"{name}.md").write_text(first_line + "\n正文\n", encoding="utf-8")

    marked = mark_navigation_units(str(tmp_path / "units"), "材料")

    assert marked == 1
    assert (root / "单元1" / ".skip_proofread").is_file()
    assert not (root / "单元2" / ".skip_proofread").exists()


def test_structural_split_rejects_out_of_range_block(tmp_path, monkeypatch):
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path))
    raw = tmp_path / "材料_raw.md"
    raw.write_text("**练1**\n题干\n", encoding="utf-8")
    blocks = [{"kind": "question", "start_line": 99, "end_line": 99, "title": "**练1**"}]

    with pytest.raises(ContractError):
        split_structural(str(raw), blocks, str(tmp_path / "units"), "材料")
