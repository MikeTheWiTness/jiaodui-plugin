"""工作区与材料契约：预期来自 ADR 0002，不重算实现逻辑。"""
import json
import shutil
from pathlib import Path

import pytest

from jiaodui.errors import JiaoduiError


def test_write_requires_workspace(tmp_path, monkeypatch):
    from jiaodui.convert import convert_to_raw
    monkeypatch.delenv("JIAODUI_WORK_ROOT", raising=False)
    src = tmp_path / "外部.md"
    src.write_text("1．题干。", encoding="utf-8")
    with pytest.raises(JiaoduiError) as err:
        convert_to_raw(str(src), str(tmp_path / "raw"), "外部")
    assert err.value.code == "workspace-required"


def test_symlink_destination_refused(tmp_path, monkeypatch):
    from jiaodui.convert import convert_to_raw
    workspace = tmp_path / "W"
    workspace.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (workspace / "raw").symlink_to(outside, target_is_directory=True)
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    src = tmp_path / "材料.md"
    src.write_text("1．题干。", encoding="utf-8")
    with pytest.raises(JiaoduiError) as err:
        convert_to_raw(str(src), str(workspace / "raw"), "材料")
    assert err.value.code == "outside-workspace"
    assert list(outside.iterdir()) == []


def _prepare(tmp_path, monkeypatch, name="A"):
    from jiaodui.workflow import convert_material, split_material
    workspace = tmp_path / name
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    src = tmp_path / "材料.md"
    src.write_text("1．设a为O点。\n\n2．第二题。\n", encoding="utf-8")
    converted = convert_material(src)
    result = split_material(converted.raw_md, subject="高中物理", mode="exam")
    return workspace, src, converted, result


def _deliver(unit):
    from jiaodui.workflow import parse_unit
    source = (unit / f"{unit.name}.md").read_text(encoding="utf-8")
    (unit / "_校对报告.md").write_text(
        f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n", encoding="utf-8")
    parse_unit(unit)


def test_layout_resume_rerun_and_move(tmp_path, monkeypatch):
    from jiaodui.workflow import split_material
    from jiaodui.status import scan_status
    from jiaodui.report import build_report
    workspace, src, converted, split = _prepare(tmp_path, monkeypatch)
    material = workspace / "校对" / "材料"
    assert Path(converted.raw_md) == material / "raw" / "材料_raw.md"
    assert (material / "source" / "材料.md").read_bytes() == src.read_bytes()
    assert len(split.unit_dirs) == 2
    paper = material / "units" / "材料"
    for unit in map(Path, split.unit_dirs):
        _deliver(unit)
    before = (paper / "第1题" / "_校对报告.md").read_bytes()
    split_material(converted.raw_md, subject="高中物理", mode="exam")
    assert scan_status(paper).counts()["已完成"] == 2
    assert (paper / "第1题" / "_校对报告.md").read_bytes() == before
    assert Path(build_report(paper).out_path).parent == material / "校对报告"
    split_material(converted.raw_md, subject="高中物理", mode="exam", rerun=True)
    assert scan_status(paper).counts().get("本轮未完成", 0) == 2
    assert not build_report(paper).included
    for unit in map(Path, split.unit_dirs):
        _deliver(unit)
    moved = tmp_path / "B" / "校对" / "材料"
    moved.parent.mkdir(parents=True)
    shutil.move(material, moved)
    monkeypatch.delenv("JIAODUI_WORK_ROOT", raising=False)
    assert scan_status(moved / "units" / "材料").counts()["已完成"] == 2
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path / "B"))
    assert Path(build_report(moved / "units" / "材料").out_path).parent == moved / "校对报告"


def test_conflict_keeps_existing_bytes(tmp_path, monkeypatch):
    from jiaodui.workflow import convert_material
    workspace, src, converted, _ = _prepare(tmp_path, monkeypatch)
    manifest = workspace / "校对" / "材料" / "_校对记录.json"
    before = manifest.read_bytes(), Path(converted.raw_md).read_bytes()
    src.write_text("源文已改变", encoding="utf-8")
    with pytest.raises(JiaoduiError) as err:
        convert_material(src)
    assert err.value.code == "material-conflict"
    assert (manifest.read_bytes(), Path(converted.raw_md).read_bytes()) == before


@pytest.mark.parametrize("kind", ["extra", "missing"])
def test_unit_set_mismatch_is_reported_not_raised(tmp_path, monkeypatch, kind):
    """R-D04：集合不一致时只读 status **不拒绝**（诊断优先），完整现场作为告警带出来。"""
    from jiaodui.status import scan_status

    workspace, _, _, _ = _prepare(tmp_path, monkeypatch)
    paper = workspace / "校对" / "材料" / "units" / "材料"
    if kind == "extra":
        (paper / "第3题").mkdir()
    else:
        shutil.rmtree(paper / "第2题")

    result = scan_status(paper)  # 不抛
    assert result.unit_set_mismatch["extra_units"] == (["第3题"] if kind == "extra" else [])
    assert result.unit_set_mismatch["missing_units"] == (["第2题"] if kind == "missing" else [])
    assert "--replace-units" in result.unit_set_mismatch["hint"]
    assert any("不一致" in w for w in result.warnings)
    # 仍然继续扫描磁盘上的真实单元
    assert result.counts()["total"] == (3 if kind == "extra" else 1)


def test_status_cli_exits_zero_on_unit_set_mismatch(tmp_path, monkeypatch):
    """R-D04 的对外表现：status 退 0，不一致在 warning 与 unit_set_mismatch 里。"""
    import os
    import subprocess
    import sys

    workspace, _, _, _ = _prepare(tmp_path, monkeypatch)
    paper = workspace / "校对" / "材料" / "units" / "材料"
    (paper / "第3题").mkdir()

    proc = subprocess.run(
        [sys.executable, "-m", "jiaodui", "status", str(paper), "--json",
         "--work-root", str(workspace)],
        capture_output=True, text=True, cwd=Path(__file__).resolve().parent.parent,
        env={**os.environ})
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["unit_set_mismatch"]["extra_units"] == ["第3题"]
    assert any("不一致" in w for w in payload["warnings"])


def test_scanning_other_paper_dir_is_still_rejected(tmp_path, monkeypatch):
    """集合不一致可以容忍，但把「另一个已存在的单元目录」当卷扫仍应拒绝。"""
    from jiaodui.status import scan_status

    workspace, _, _, _ = _prepare(tmp_path, monkeypatch)
    other = workspace / "校对" / "材料" / "units" / "材料" / "第1题"
    with pytest.raises(JiaoduiError) as err:
        scan_status(other)
    assert err.value.code == "paper-path-mismatch"


def test_corrupt_manifest_never_legacy_fallback(tmp_path, monkeypatch):
    from jiaodui.status import scan_status
    workspace, _, _, _ = _prepare(tmp_path, monkeypatch)
    material = workspace / "校对" / "材料"
    (material / "_校对记录.json").write_text("{坏JSON", encoding="utf-8")
    with pytest.raises(JiaoduiError) as err:
        scan_status(material / "units" / "材料")
    assert err.value.code == "invalid-manifest"


def _png_bytes(color) -> bytes:
    """真实 PNG 字节：图片已是闸门项，假字节会让单元无法交付。"""
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(buffer, "PNG")
    return buffer.getvalue()


def test_changed_markdown_image_starts_new_run(tmp_path, monkeypatch):
    from jiaodui.workflow import convert_material, split_material
    from jiaodui.status import scan_status
    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source = tmp_path / "卷.md"
    source.write_text("1．题干 ![](pic.png)。\n", encoding="utf-8")
    image = tmp_path / "pic.png"
    second = _png_bytes("blue")
    image.write_bytes(_png_bytes("red"))
    converted = convert_material(source)
    split = split_material(converted.raw_md, subject="高中物理", mode="exam")
    unit = Path(split.unit_dirs[0])
    _deliver(unit)
    paper = unit.parent
    manifest = workspace / "校对" / "卷" / "_校对记录.json"
    before = json.loads(manifest.read_text())["current_run_id"]
    image.write_bytes(second)
    convert_material(source)
    assert scan_status(paper).counts()["已完成"] == 0
    split_material(converted.raw_md, subject="高中物理", mode="exam")
    assert json.loads(manifest.read_text())["current_run_id"] != before
    assert (unit / "images" / "pic.png").read_bytes() == second


def test_missing_external_image_cannot_reuse_previous_raw_media(tmp_path, monkeypatch):
    """重转外部材料必须反映源图缺失，使本轮完成失效；显式 source 副本仍可独立复用。"""
    from jiaodui import material
    from jiaodui.status import scan_status
    from jiaodui.workflow import convert_material, split_material

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source_dir = tmp_path / "external"
    source_dir.mkdir()
    source = source_dir / "卷.md"
    source.write_text("1．题干 ![](pic.png)。\n", encoding="utf-8")
    image = source_dir / "pic.png"
    original_image = _png_bytes("red")
    image.write_bytes(original_image)
    converted = convert_material(source, mode="exam")
    split = split_material(converted.raw_md, subject="高中物理", mode="exam")
    unit = Path(split.unit_dirs[0])
    _deliver(unit)
    root = workspace / "校对" / "卷"
    before_run = material.load(root)["current_run_id"]
    assert scan_status(unit.parent).counts()["已完成"] == 1

    assert image.resolve() == (source_dir / "pic.png").resolve()
    image.unlink()
    result = convert_material(source, mode="exam")

    assert (result.copied, result.missing) == (0, 1)
    assert any("图片未找到" in warning for warning in result.warnings)
    assert "![](pic.png)" in Path(result.raw_md).read_text(encoding="utf-8")
    assert not (Path(result.images_dir) / "media" / "pic.png").exists()
    assert material.load(root)["current_run_id"] != before_run
    assert scan_status(unit.parent).counts()["已完成"] == 0

    # 源副本保留导入时的资源；调用方明确选择该副本时，允许按它自己的相对路径重转。
    saved_source = root / "source" / "卷.md"
    assert (saved_source.parent / "pic.png").read_bytes() == original_image
    restored = convert_material(saved_source, mode="exam")
    assert (restored.copied, restored.missing) == (1, 0), restored.warnings
    assert (Path(restored.images_dir) / "media" / "pic.png").read_bytes() == original_image


def test_adopt_missing_units_is_recorded(tmp_path, monkeypatch):
    from jiaodui.workflow import split_material
    from jiaodui.status import scan_status
    workspace, _, converted, result = _prepare(tmp_path, monkeypatch)
    paper = Path(result.unit_dirs[0]).parent
    shutil.rmtree(paper / "第2题")
    split_material(converted.raw_md, subject="高中物理", mode="exam", adopt_units=True)
    assert scan_status(paper).counts()["total"] == 1
    data = json.loads((workspace / "校对" / "材料" / "_校对记录.json").read_text())
    action = data["runs"][-1]["actions"][-1]
    assert action["before"] == ["第1题", "第2题"]
    assert action["after"] == ["第1题"]


def test_failed_resplit_does_not_wedge_material(tmp_path, monkeypatch):
    """拆分结果与当前集合不一致时：不写清单，材料仍自洽可用。

    旧实现会先写清单（unit_set=新集合）再核对磁盘，于是磁盘留着旧目录、
    清单已是新集合，之后连只读的 status 都会报 unit-set-mismatch。
    """
    from jiaodui.status import scan_status
    from jiaodui.workflow import split_material

    _, _, converted, result = _prepare(tmp_path, monkeypatch)  # 第1题 / 第2题
    paper = Path(result.unit_dirs[0]).parent

    with pytest.raises(JiaoduiError) as err:
        split_material(converted.raw_md, subject="高中物理", mode="exam", strategy="none")
    assert err.value.code == "unit-set-mismatch"

    assert scan_status(paper).counts()["total"] == 2


@pytest.mark.parametrize("strategy", ["rule", "none", "manual", "boundaries"])
def test_rejected_resplit_preserves_completed_material(tmp_path, monkeypatch, strategy):
    """集合核对必须先于写盘：被拒绝的重拆不能改坏旧源文、图片或交付登记。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import split_material

    workspace, _, converted, result = _prepare(tmp_path, monkeypatch)
    units = list(map(Path, result.unit_dirs))
    paper = units[0].parent
    for unit in units:
        _deliver(unit)
    assert scan_status(paper).counts()["已完成"] == 2

    raw = Path(converted.raw_md)
    args = {"strategy": strategy}
    if strategy == "rule":
        raw.write_text("1．新的第一题。\n", encoding="utf-8")
    elif strategy == "manual":
        raw.write_text(
            "###### 单元开始 ######\n新的第一题。\n###### 单元结束 ######\n",
            encoding="utf-8")
    elif strategy == "boundaries":
        args = {"scope": "excerpt", "scope_reason": "人工集合变更测试", "boundaries": [
            {"name": "第2题", "text": "被拒绝的新内容 ![](pic.png)"},
            {"name": "第3题", "text": "新的第三题。"},
        ]}
        media = Path(converted.images_dir) / "media"
        (media / "pic.png").write_bytes(_png_bytes("blue"))
        # 若先写后拒，这张旧单元图片也会被覆盖。
        (units[1] / "images" / "pic.png").write_bytes(_png_bytes("red"))
        split_material(converted.raw_md, subject="高中物理", mode="exam", adopt_units=True)
        for unit in units:
            _deliver(unit)

    material_dir = workspace / "校对" / "材料"
    before = {p.relative_to(material_dir): p.read_bytes()
              for p in material_dir.rglob("*") if p.is_file()}
    with pytest.raises(JiaoduiError) as err:
        split_material(converted.raw_md, subject="高中物理", mode="exam", **args)

    assert err.value.code == "unit-set-mismatch"
    after = {p.relative_to(material_dir): p.read_bytes()
             for p in material_dir.rglob("*") if p.is_file()}
    assert after == before, "拒绝拆分必须发生在任何单元写盘之前"
    assert scan_status(paper).counts()["已完成"] == 2
    assert not (paper / "第3题").exists()


@pytest.mark.parametrize("replace_units", [False, True])
def test_duplicate_slice_names_preserve_completed_material(tmp_path, monkeypatch, replace_units):
    """重复边界（包括允许替换集合）必须先拒绝，材料全部字节与完成状态不变。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import convert_material, split_material

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source_dir = tmp_path / "external"
    source_dir.mkdir()
    source = source_dir / "讲义.md"
    source.write_text("第一段 ![](pic.png)\n第二段\n第三段", encoding="utf-8")
    (source_dir / "pic.png").write_bytes(_png_bytes("red"))
    converted = convert_material(source, mode="lecture")
    result = split_material(converted.raw_md, subject="高中物理", mode="lecture", strategy="none")
    unit = Path(result.unit_dirs[0])
    _deliver(unit)
    # 如果预检发生在写盘后，该蓝图会覆盖原单元红图。
    (Path(converted.images_dir) / "media" / "pic.png").write_bytes(_png_bytes("blue"))
    root = workspace / "校对" / "讲义"
    before = {path.relative_to(root): path.read_bytes()
              for path in root.rglob("*") if path.is_file()}

    with pytest.raises(JiaoduiError) as err:
        split_material(converted.raw_md, subject="高中物理", mode="lecture",
                       replace_units=replace_units, boundaries=[
                           {"name": "单元1", "start_line": 1, "end_line": 1},
                           {"name": "单元1", "start_line": 3, "end_line": 3},
                       ])

    after = {path.relative_to(root): path.read_bytes()
             for path in root.rglob("*") if path.is_file()}
    assert after == before, "拒绝重复名称必须早于源文、图片、清单和归档写入"
    assert err.value.code == "duplicate-unit-name"
    assert err.value.details["duplicate_names"] == ["单元1"]
    assert scan_status(unit.parent).counts()["已完成"] == 1
    assert not (unit.parent / ".superseded").exists()


@pytest.mark.parametrize("replacement", ["**例1**\n新的例题正文。\n", ""])
def test_rejected_lecture_resplit_keeps_sources(tmp_path, monkeypatch, replacement):
    """讲义规则拆分与空结果同样必须在覆盖旧单元之前拒绝。"""
    from jiaodui.workflow import convert_material, split_material

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source = tmp_path / "讲义.md"
    source.write_text("**例1**\n例题甲正文。\n\n**例2**\n例题乙正文。\n", encoding="utf-8")
    converted = convert_material(source, mode="lecture")
    result = split_material(converted.raw_md, subject="高中物理", mode="lecture")
    assert [Path(p).name for p in result.unit_dirs] == ["单元1", "单元2"]
    sources = [Path(p) / f"{Path(p).name}.md" for p in result.unit_dirs]
    before = [p.read_bytes() for p in sources]
    Path(converted.raw_md).write_text(replacement, encoding="utf-8")

    with pytest.raises(JiaoduiError) as err:
        split_material(converted.raw_md, subject="高中物理", mode="lecture")

    assert err.value.code == "unit-set-mismatch"
    assert [p.read_bytes() for p in sources] == before


def test_resplit_can_still_add_units_without_replacement(tmp_path, monkeypatch):
    """保留既有增补口径：没有淘汰旧单元时不要求 --replace-units。"""
    from jiaodui.workflow import split_material

    _, _, converted, result = _prepare(tmp_path, monkeypatch)
    units = list(map(Path, result.unit_dirs))
    before = [(unit / f"{unit.name}.md").read_bytes() for unit in units]
    added = split_material(converted.raw_md, subject="高中物理", mode="exam",
                           scope="excerpt", scope_reason="人工增补正文测试",
                           boundaries=[
                               {"name": unit.name, "text": data.decode("utf-8")}
                               for unit, data in zip(units, before)
                           ] + [{"name": "第3题", "text": "3．第三题。"}])

    assert [Path(p).name for p in added.unit_dirs] == ["第1题", "第2题", "第3题"]
    assert [(unit / f"{unit.name}.md").read_bytes() for unit in units] == before
    assert (units[0].parent / "第3题" / "第3题.md").read_text(encoding="utf-8") == "3．第三题。"


def test_failed_resplit_does_not_create_new_units(tmp_path, monkeypatch):
    """集合不一致时不得留下本轮新增目录（先核对则根本不创建）。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import slice_material

    _, _, converted, result = _prepare(tmp_path, monkeypatch)
    paper = Path(result.unit_dirs[0]).parent
    raw = Path(converted.raw_md)
    spec = tmp_path / "b.json"
    # 新集合淘汰第1题并计划新增第3题，必须在修改第2题或创建第3题前拒绝。
    spec.write_text(json.dumps({
        "raw": str(raw), "mode": "exam",
        "boundaries": [
            {"name": "第2题", "start_line": 3, "end_line": 3},
            {"name": "第3题", "start_line": 1, "end_line": 1},
        ],
    }, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(JiaoduiError) as err:
        slice_material(spec, subject="高中物理")
    assert err.value.code == "unit-set-mismatch"
    assert not (paper / "第3题").exists(), "被拒绝的拆分不得创建新目录"
    assert scan_status(paper).counts()["total"] == 2


def test_md_source_does_not_absorb_files_outside_its_directory(tmp_path, monkeypatch):
    """源目录之外的引用不得被复制进材料：否则一份构造的 md 可吸走宿主机任意可读文件。"""
    from jiaodui.workflow import convert_material

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    secret_dir = tmp_path / "外部"
    secret_dir.mkdir()
    secret = secret_dir / "secret.txt"
    secret.write_text("TOP SECRET", encoding="utf-8")
    # 源 md 放在自己的子目录里，secret 才是真正的「源目录之外」
    md_dir = tmp_path / "材料目录"
    md_dir.mkdir()
    md = md_dir / "材料.md"
    md.write_text(f"1．题干 ![]({secret})\n", encoding="utf-8")

    convert_material(md)

    absorbed = [p for p in workspace.rglob("*") if p.is_file() and p.name.endswith("secret.txt")]
    assert not absorbed, f"不该把源目录之外的引用复制进材料：{absorbed}"


def test_unit_set_mismatch_points_at_replace_units(tmp_path, monkeypatch):
    """集合不一致时要给出可执行的恢复指令，而不是只有一句「不一致」。"""
    from jiaodui.workflow import split_material

    _, _, converted, result = _prepare(tmp_path, monkeypatch)
    paper = Path(result.unit_dirs[0]).parent
    (paper / "第3题").mkdir()
    with pytest.raises(JiaoduiError) as err:
        split_material(converted.raw_md, subject="高中物理", mode="exam")
    assert err.value.code == "unit-set-mismatch"
    details = json.dumps(err.value.details, ensure_ascii=False)
    assert "--replace-units" in details
    assert "第3题" in details


def test_replace_units_archives_stale_and_keeps_history(tmp_path, monkeypatch):
    """换单元集：陈旧单元归档到 .superseded/<run_id>/，只移动不删除，并写进轮次记录。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import split_material

    workspace, _, converted, result = _prepare(tmp_path, monkeypatch)  # 第1题 / 第2题
    paper = Path(result.unit_dirs[0]).parent
    manifest = workspace / "校对" / "材料" / "_校对记录.json"
    run_id = json.loads(manifest.read_text())["current_run_id"]
    # 给被淘汰的单元留一份报告，验证归档后仍可追溯
    stale = paper / "第2题"
    (stale / "_校对报告.md").write_text("无问题\n### 标记原文\n2．第二题。\n### 修改原因\n无\n",
                                        encoding="utf-8")

    replaced = split_material(converted.raw_md, subject="高中物理", mode="exam",
                              strategy="none", replace_units=True)

    assert [Path(p).name for p in replaced.unit_dirs] == ["第1题"]
    assert not stale.exists(), "陈旧单元应从单元根目录移走"
    archived = paper / ".superseded" / run_id / "第2题"
    assert archived.is_dir()
    assert (archived / "_校对报告.md").is_file(), "历史报告必须保留"
    assert (archived / "第2题.md").is_file(), "单元源文必须保留"

    assert scan_status(paper).counts()["total"] == 1
    data = json.loads(manifest.read_text())
    assert data["unit_set"] == ["第1题"]
    moved = data["runs"][-1]["actions"][-1]["args"]["replace_units"]
    assert moved == [{"unit": "第2题",
                      "from": "材料/第2题",
                      "to": f"材料/.superseded/{run_id}/第2题"}]


def test_param_change_preserves_report_but_invalidates_registration(tmp_path, monkeypatch):
    from jiaodui.workflow import split_material
    from jiaodui.status import scan_status
    _, _, converted, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    _deliver(unit)
    before = (unit / "_校对报告.md").read_bytes()
    split_material(converted.raw_md, subject="高中物理", mode="exam", clean=False)
    assert (unit / "_校对报告.md").read_bytes() == before
    assert scan_status(unit.parent).counts()["已完成"] == 0


def test_manifest_escape_without_workspace_is_rejected(tmp_path, monkeypatch):
    from jiaodui.status import scan_status
    workspace, _, _, result = _prepare(tmp_path, monkeypatch)
    manifest = workspace / "校对" / "材料" / "_校对记录.json"
    data = json.loads(manifest.read_text())
    data["paths"]["raw"] = "../../outside.md"
    manifest.write_text(json.dumps(data), encoding="utf-8")
    monkeypatch.delenv("JIAODUI_WORK_ROOT", raising=False)
    with pytest.raises(JiaoduiError) as err:
        scan_status(Path(result.unit_dirs[0]).parent)
    assert err.value.code == "invalid-manifest"


def test_docx_raw_images_remain_relative_after_move(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    from jiaodui.convert import find_pandoc
    from jiaodui.workflow import convert_material, split_material
    if not find_pandoc():
        pytest.skip("需要 pandoc")
    image = tmp_path / "picture.png"
    Image.new("RGB", (20, 20), "red").save(image)
    source = tmp_path / "图卷.docx"
    document = Document()
    document.add_paragraph("1．看图判断。")
    document.add_picture(str(image))
    document.save(source)
    workspace = tmp_path / "A"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    converted = convert_material(source)
    assert converted.missing == 0 and converted.copied == 1
    text = Path(converted.raw_md).read_text()
    assert "./图卷_images/media/" in text
    moved = tmp_path / "B" / "校对" / "图卷"
    moved.parent.mkdir(parents=True)
    shutil.move(workspace / "校对" / "图卷", moved)
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(tmp_path / "B"))
    split = split_material(moved / "raw" / "图卷_raw.md", subject="高中物理", mode="exam")
    assert split.missing == 0 and split.copied == 1


def test_subprocess_outputs_rejected_before_spawn(tmp_path, monkeypatch):
    from jiaodui import convert
    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    def forbidden(*args, **kwargs):
        pytest.fail("越界输出不能启动子进程")
    monkeypatch.setattr(convert.subprocess, "run", forbidden)
    with pytest.raises(JiaoduiError) as err:
        convert.convert_with_pandoc("input.docx", str(tmp_path / "outside.md"), str(workspace / "imgs"))
    assert err.value.code == "outside-workspace"


def test_concurrent_registration_keeps_all_units(tmp_path, monkeypatch):
    import subprocess
    import sys
    from jiaodui.status import scan_status
    from jiaodui.workflow import split_material, convert_material
    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source = tmp_path / "并发.md"
    source.write_text("\n\n".join(f"{n}．题干。" for n in range(1, 9)), encoding="utf-8")
    converted = convert_material(source)
    split = split_material(converted.raw_md, subject="高中物理", mode="exam")
    for d in map(Path, split.unit_dirs):
        src = (d / f"{d.name}.md").read_text()
        (d / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{src}\n### 修改原因\n无\n")
    jobs = [subprocess.Popen([sys.executable, "-m", "jiaodui", "parse-report", "--unit", d, "--json"],
                             stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True) for d in split.unit_dirs]
    for job in jobs:
        stdout, stderr = job.communicate(timeout=20)
        assert job.returncode == 0, stderr
        assert json.loads(stdout)["registered"]
    assert scan_status(Path(split.unit_dirs[0]).parent).counts()["已完成"] == 8


def test_word_rejects_partial_current_run(tmp_path, monkeypatch):
    from jiaodui.docx_report import build_docx, find_pandoc
    if not find_pandoc():
        pytest.skip("需要 pandoc")
    _, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    _deliver(unit)
    word = build_docx(str(unit.parent))
    assert not word.ok
    assert [u["unit"] for u in word.excluded_units] == ["第2题"]


def test_material_name_keeps_literal_source_stem(tmp_path, monkeypatch):
    from jiaodui.workflow import convert_material
    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source = tmp_path / "源_raw.md"
    source.write_text("1．题干。")
    converted = convert_material(source)
    assert Path(converted.raw_md) == workspace / "校对" / "源_raw" / "raw" / "源_raw_raw.md"


def test_failed_conversion_can_resume_same_material(tmp_path, monkeypatch):
    from jiaodui import convert
    from jiaodui.workflow import convert_material
    from jiaodui.errors import EnvError
    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    source = tmp_path / "失败.md"
    source.write_text("1．题干。")
    original = convert.convert_to_raw
    def fail(*args, **kwargs):
        raise EnvError("模拟转换失败")
    monkeypatch.setattr(convert, "convert_to_raw", fail)
    with pytest.raises(EnvError):
        convert_material(source)
    monkeypatch.setattr(convert, "convert_to_raw", original)
    assert Path(convert_material(source).raw_md).is_file()


def test_docx_reads_free_form_source_like_status_does(tmp_path):
    """R-D02：单元源文只有一个定义。

    把 `第1题.md` 改名成 `草稿副本.md` 后，status 认它（find_source_md 回退），
    docx 的标记审计也必须认它——旧实现只认契约名，会静默读不到源文、跳过比对。
    """
    from jiaodui import paths
    from jiaodui.docx_report import _read_unit_source
    from jiaodui.status import status_of_unit

    unit = tmp_path / "第1题"
    unit.mkdir()
    (unit / "草稿副本.md").write_text("设a为$O$点。", encoding="utf-8")
    (unit / "_校对报告.md").write_text(
        "无问题\n### 标记原文\n设a为$O$点。\n### 修改原因\n无\n", encoding="utf-8")

    assert paths.find_source_md(unit).name == "草稿副本.md"
    assert status_of_unit(unit).state == "已完成"
    assert _read_unit_source(tmp_path, "第1题") == "设a为$O$点。"


def test_extra_unit_files_are_tolerated_and_inert(tmp_path, monkeypatch):
    """R-D03（决策：不管）：非契约文件不导致失败，也不参与校验与交付。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import parse_unit

    _, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    (unit / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n")
    (unit / "第1题_尝试2.md").write_text("旧草稿", encoding="utf-8")
    (unit / "随手笔记.txt").write_text("笔记", encoding="utf-8")

    assert parse_unit(unit)["registered"]
    assert scan_status(unit.parent).counts()["已完成"] == 1


def test_readonly_tolerates_symlinked_source_but_write_rejects(tmp_path, monkeypatch):
    """R-D07（决策）：只读入口允许软链越界，写入口仍按 realpath 拒绝。"""
    from jiaodui.status import scan_status
    from jiaodui.verify import verify_unit
    from jiaodui.workflow import parse_unit

    workspace = tmp_path / "W"
    workspace.mkdir()
    monkeypatch.setenv("JIAODUI_WORK_ROOT", str(workspace))
    outside = tmp_path / "外" / "第1题.md"
    outside.parent.mkdir()
    outside.write_text("设a为$O$点。", encoding="utf-8")

    unit = workspace / "第1题"
    unit.mkdir()
    (unit / "第1题.md").symlink_to(outside)
    (unit / "_校对报告.md").write_text(
        "无问题\n### 标记原文\n设a为$O$点。\n### 修改原因\n无\n", encoding="utf-8")

    # 只读：读得到、不报边界错
    assert verify_unit(unit).ok
    scan_status(workspace)  # 不抛即达标

    # 写：同一单元按 realpath 拒绝，绝不顺着软链写出去
    with pytest.raises(JiaoduiError) as err:
        parse_unit(unit)
    assert err.value.code == "outside-workspace"


def test_build_report_out_must_stay_in_contract_dir(tmp_path, monkeypatch):
    """R-D09：有清单时 --out 必须落在契约目录 校对报告/ 内，文件名可自取。"""
    from jiaodui.report import build_report

    workspace, _, _, result = _prepare(tmp_path, monkeypatch)
    paper = Path(result.unit_dirs[0]).parent

    with pytest.raises(JiaoduiError) as err:
        build_report(paper, out_path=workspace / "随便.md")
    assert err.value.code == "out-of-contract"

    inside = workspace / "校对" / "材料" / "校对报告" / "自定义名.md"
    built = build_report(paper, out_path=inside)
    assert built.out_path == str(inside)
    assert inside.is_file()


def test_build_docx_out_dir_must_stay_in_contract_dir(tmp_path, monkeypatch):
    """R-D09：--out-dir 同样收口，且在建目录之前就拒绝（不留垃圾目录）。"""
    from jiaodui.docx_report import build_docx

    workspace, _, _, result = _prepare(tmp_path, monkeypatch)
    paper = Path(result.unit_dirs[0]).parent
    outside = workspace / "随便目录"

    with pytest.raises(JiaoduiError) as err:
        build_docx(str(paper), out_dir=str(outside))
    assert err.value.code == "out-of-contract"
    assert not outside.exists(), "拒绝时不得留下目录"


def test_missing_data_file_breaks_current_run_completion(tmp_path, monkeypatch):
    """R-D06：_校对数据.json 是交付的一部分，删掉就不再算本轮完成。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import parse_unit

    _, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    (unit / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n")
    assert parse_unit(unit)["registered"]
    assert scan_status(unit.parent).counts()["已完成"] == 1

    (unit / "_校对数据.json").unlink()

    counts = scan_status(unit.parent).counts()
    assert counts["已完成"] == 0
    assert counts["本轮未完成"] == 1


def test_tampered_data_file_breaks_current_run_completion(tmp_path, monkeypatch):
    """改过 _校对数据.json 也不能算本轮完成（与登记摘要不一致）。"""
    from jiaodui.status import scan_status
    from jiaodui.workflow import parse_unit

    _, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    (unit / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n")
    parse_unit(unit)

    data_file = unit / "_校对数据.json"
    data_file.write_text(data_file.read_text(encoding="utf-8") + "\n", encoding="utf-8")

    counts = scan_status(unit.parent).counts()
    assert counts["已完成"] == 0
    assert counts["本轮未完成"] == 1


def test_registration_records_data_digest(tmp_path, monkeypatch):
    """登记必须把数据文件摘要写进 unit_records，否则上面两条无从判定。"""
    from jiaodui.workflow import parse_unit

    workspace, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    (unit / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n")
    parse_unit(unit)

    manifest = workspace / "校对" / "材料" / "_校对记录.json"
    records = json.loads(manifest.read_text(encoding="utf-8"))["unit_records"]
    assert len(records) == 1
    assert len(records[0]["data_sha256"]) == 64


def test_registration_refuses_run_changed_during_parse(tmp_path, monkeypatch):
    from jiaodui.workflow import parse_unit, split_material
    from jiaodui import report_parse
    _, _, converted, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    (unit / "_校对报告.md").write_text(f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n")
    save = report_parse.save_proofread_json
    def new_run_before_register(*args, **kwargs):
        saved = save(*args, **kwargs)
        split_material(converted.raw_md, subject="高中物理", mode="exam", rerun=True)
        return saved
    monkeypatch.setattr(report_parse, "save_proofread_json", new_run_before_register)
    with pytest.raises(JiaoduiError) as err:
        parse_unit(unit)
    assert err.value.code == "run-changed"


def test_crlf_report_registers_its_actual_bytes(tmp_path, monkeypatch):
    from jiaodui.workflow import parse_unit
    from jiaodui.status import scan_status
    _, _, _, result = _prepare(tmp_path, monkeypatch)
    unit = Path(result.unit_dirs[0])
    source = (unit / "第1题.md").read_text()
    report = f"无问题\n### 标记原文\n{source}\n### 修改原因\n无\n".replace("\n", "\r\n")
    (unit / "_校对报告.md").write_bytes(report.encode("utf-8"))
    assert parse_unit(unit)["registered"]
    assert scan_status(unit.parent).counts()["已完成"] == 1
